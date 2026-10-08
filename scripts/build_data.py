#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Kablaj Senelik Özetler -> data/summary.json üretir.

Kullanım:
    python scripts/build_data.py                      # data/*.xlsx -> data/summary.json
    python scripts/build_data.py --kaynak "C:/Users/user/Desktop/Kablaj Senelik Özetler" --hafif

--hafif bayrağı: satır detayını JSON'a gömmez (dosya küçük olur, site filtreleri
kısıtlı çalışır). Varsayılan: satırlar dahil (filtreler tam çalışır, dosya ~3-6 MB).
GitHub Pages 25 MB altı dosyaları sorunsuz servis eder.

Beklenen girdi: SRM dökümü, Sheet1 ilk satır başlık. Kritik kolonlar isimden bulunur:
  Üretim Yeri / Satıcı Adı / Malzeme No / Malzeme Tanımı / Denetim Sonucu / Yaratılma Tarihi
"""
import argparse, datetime, json, re, sys, unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl gerekli: pip install openpyxl")

AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Mayyıs", "Haziran",
         "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
AY_MAP = {"Mayyıs": "Mayıs", "Mayis": "Mayıs", "Subat": "Şubat", "Mart ": "Mart"}

KOK = Path(__file__).resolve().parent.parent
VARSAYILAN_KAYNAK = KOK / "data"

def norm(s):
    # Türkçe uyumlu: NFKD ile aksanları sök, İ->I->i ve ı->i dönüşümlerini elle yap.
    # (NFKD tek başına 'ı' harfini dönüştürmez; bu yüzden Satıcı Adı gibi
    # kolonlar eşleşemezdi.)
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().replace("ı", "i")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()

# SRM dökümünün bilinen sabit kolon sırası (isim eşleşmesi başarısız olursa yedek)
YEDEK_INDEX = {"uretim": 3, "satici": 7, "malz_no": 8, "malz_ad": 11,
               "sonuc": 31, "kontrol": 4, "tarih": 22, "denetim_turu": 14}

def kolon_bul(basliklar):
    n = [norm(b) for b in basliklar]
    def bul(*adaylar):
        for a in adaylar:
            na = norm(a)
            for i, h in enumerate(n):
                if na in h:
                    return i
        return -1
    return {
        "uretim": bul("uretim yeri"),
        "satici": bul("satici adi"),
        "malz_no": bul("malzeme no"),
        "malz_ad": bul("malzeme tanimi"),
        "sonuc": bul("denetim sonucu"),
        "kontrol": bul("kontrol turu"),
        "tarih": bul("yaratilm"),
        "denetim_turu": bul("denetim turu"),
    }

def sonuc_sinif(v):
    t = norm(v)
    if not t:
        return "bos"
    if "degil" in t or t.startswith("red") or "ret" in t:
        return "red"
    if "uygun" in t:
        return "uygun"
    return "diger"

def ay_etiketi(dosya_adi):
    govde = Path(dosya_adi).stem.strip()
    for a in AYLAR:
        if norm(a) in norm(govde):
            return AY_MAP.get(a, a)
    return govde

def dosyayi_oku(yol, ay):
    wb = openpyxl.load_workbook(str(yol), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    satirlar = list(ws.iter_rows(values_only=True))
    if len(satirlar) < 2:
        return []
    baslik = [str(c or "") for c in satirlar[0]]
    K = kolon_bul(baslik)
    for k, v in YEDEK_INDEX.items():  # isimle bulunamayanı sabit sıradan al
        if K[k] < 0 and v < len(baslik):
            K[k] = v
    if K["sonuc"] < 0:
        print(f"  UYARI: {yol.name} içinde 'Denetim Sonucu' kolonu bulunamadı, atlanıyor.")
        return []
    kayitlar = []
    for r in satirlar[1:]:
        def al(i):
            if i < 0 or i >= len(r):
                return ""
            v = r[i]
            return "" if v is None else str(v).strip()
        sonuc = al(K["sonuc"])
        kayitlar.append({
            "ay": ay,
            "uretim": al(K["uretim"]) or "Bilinmiyor",
            "satici": al(K["satici"]) or "Bilinmiyor",
            "malzNo": al(K["malz_no"]),
            "malzAd": al(K["malz_ad"])[:80],
            "kontrol": al(K["kontrol"]) or "—",
            "tarih": al(K["tarih"]),
            "sonuc": sonuc,
            "sinif": sonuc_sinif(sonuc),
        })
    return kayitlar

def ozetle(kayitlar):
    def oran(red, top):
        return round(red / top * 100, 2) if top else 0.0
    aylar = sorted(set(k["ay"] for k in kayitlar),
                   key=lambda a: AYLAR.index(a) if a in AYLAR else 99)
    aylik = []
    for ay in aylar:
        rs = [k for k in kayitlar if k["ay"] == ay]
        red = sum(1 for k in rs if k["sinif"] == "red")
        uygun = sum(1 for k in rs if k["sinif"] == "uygun")
        aylik.append({"ay": ay, "toplam": len(rs), "uygun": uygun, "red": red,
                      "bos": len(rs) - uygun - red, "hata_oran": oran(red, len(rs))})
    return {"aylar": aylar, "aylik": aylik}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--kaynak", default=str(VARSAYILAN_KAYNAK),
                    help="xlsx klasörü (varsayılan: data/)")
    ap.add_argument("--cikti", default=str(KOK / "data" / "summary.json"))
    ap.add_argument("--hafif", action="store_true",
                    help="satır detayını gömme (küçük dosya, kısıtlı filtre)")
    args = ap.parse_args()

    kaynak = Path(args.kaynak)
    dosyalar = sorted(kaynak.glob("*.xlsx"))
    if not dosyalar:
        sys.exit(f"{kaynak} içinde xlsx bulunamadı. Aylık dosyaları data/ klasörüne kopyala.")

    tumu = []
    manifest = []
    for yol in dosyalar:
        if yol.name.startswith("~"):
            continue
        ay = ay_etiketi(yol.name)
        print(f"- {yol.name} -> {ay} okunuyor…")
        kay = dosyayi_oku(yol, ay)
        print(f"  {len(kay)} satır")
        tumu.extend(kay)
        manifest.append({"dosya": yol.name, "ay": ay})

    oz = ozetle(tumu)
    red = sum(1 for k in tumu if k["sinif"] == "red")
    uygun = sum(1 for k in tumu if k["sinif"] == "uygun")
    cikti = {
        "meta": {
            "uretilme": datetime.datetime.now().strftime("%d.%m.%Y %H:%M"),
            "aylar": oz["aylar"],
            "toplam_kayit": len(tumu),
            "uygun": uygun, "red": red,
            "hata_oran": round(red / len(tumu) * 100, 2) if tumu else 0,
            "kaynak": "build_data.py",
        },
        "aylik": oz["aylik"],
    }
    if not args.hafif:
        cikti["satirlar"] = tumu
    else:
        cikti["satirlar"] = []

    out = Path(args.cikti)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(cikti, ensure_ascii=False), encoding="utf-8")
    (KOK / "data" / "manifest.json").write_text(
        json.dumps({"dosyalar": manifest}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nOK: {out} ({out.stat().st_size/1024:.0f} KB), {len(tumu)} kayıt, "
          f"hata %{cikti['meta']['hata_oran']}")
    print("manifest.json güncellendi.")

if __name__ == "__main__":
    main()
